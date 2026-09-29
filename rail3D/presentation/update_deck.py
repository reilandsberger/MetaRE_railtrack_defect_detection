"""Add the horn + FDTD-setup (+ prelim result) slides to rail3D_overview.pptx,
in the deck's own style, and fix the text that the horn change made stale.

    pip install python-pptx          # not in requirements.txt: deck tooling only
    python presentation/update_deck.py

Run from rail3D/. ALWAYS starts from the 2026-09-04 backup
(data/generated/rail3D_overview_2026-09-04.pptx, created on first run), so it is
idempotent; figures come from data/figures/deck_2026-09-24 (deck_figs.py).
Written 2026-09-24 with minimal layout QA -- check spacing in PowerPoint.
"""
import shutil
import sys
from pathlib import Path

SP = Path(__file__).parent
from PIL import Image                                          # noqa: E402
from pptx import Presentation                                  # noqa: E402
from pptx.dml.color import RGBColor                            # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR                # noqa: E402
from pptx.util import Emu, Inches, Pt                          # noqa: E402

ROOT = Path.cwd()                                   # rail3D/
DECK = ROOT / "data/generated/rail3D_overview.pptx"
BACKUP = ROOT / "data/generated/rail3D_overview_2026-09-04.pptx"
ASSETS = ROOT / "data/figures/deck_2026-09-24"
if not BACKUP.exists():
    shutil.copy2(DECK, BACKUP)

INK = RGBColor(0x1B, 0x27, 0x33)
MUTE = RGBColor(0x7B, 0x84, 0x94)
ACC = RGBColor(0x1F, 0x6F, 0xEB)
GOOD = RGBColor(0x1A, 0x7F, 0x5A)
BAD = RGBColor(0xB4, 0x23, 0x18)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF6, 0xF8, 0xFA)

prs = Presentation(BACKUP)
BLANK = prs.slide_layouts[6]


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size=16, color=INK, bold=False, space=8, first=False,
         italic=False, align=None, font=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space)
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font or "Calibri"
    return p


def bar(slide, y=Inches(1.22)):
    s = slide.shapes.add_shape(1, Inches(0.62), y, Inches(1.05), Pt(3.5))
    s.fill.solid(); s.fill.fore_color.rgb = ACC
    s.line.fill.background(); s.shadow.inherit = False


def head(slide, kicker, title, sub=None):
    t = tb(slide, Inches(0.6), Inches(0.32), Inches(12.2), Inches(0.32))
    para(t, kicker.upper(), 11.5, MUTE, bold=True, first=True, space=0)
    t = tb(slide, Inches(0.6), Inches(0.58), Inches(12.2), Inches(0.62))
    para(t, title, 29, INK, bold=True, first=True, space=0)
    bar(slide)
    if sub:
        t = tb(slide, Inches(0.6), Inches(1.34), Inches(12.2), Inches(0.5))
        para(t, sub, 14.5, MUTE, first=True, space=0)


def picture(slide, path, y, h, x=Inches(0.6), w=Inches(12.13)):
    iw, ih = Image.open(path).size
    sc = min(int(w) / iw, int(h) / ih)
    pw, ph = int(iw * sc), int(ih * sc)
    slide.shapes.add_picture(str(path), Emu(int(x) + (int(w) - pw) // 2),
                             Emu(int(y) + (int(h) - ph) // 2), Emu(pw), Emu(ph))


def note(slide, text, color=MUTE, y=Inches(6.86), bold=False, size=13):
    t = tb(slide, Inches(0.6), y, Inches(12.13), Inches(0.5))
    para(t, text, size, color, bold=bold, italic=not bold, first=True, space=0)


def stat_row(slide, items, y, h=Inches(1.0), x0=Inches(0.6), wtot=Inches(12.13)):
    n = len(items)
    gap = Inches(0.2)
    w = int((int(wtot) - int(gap) * (n - 1)) / n)
    for i, (big, small, col) in enumerate(items):
        s = slide.shapes.add_shape(5, Emu(int(x0) + i * (w + int(gap))), y, Emu(w), h)
        s.fill.solid(); s.fill.fore_color.rgb = PANEL
        s.line.fill.background(); s.shadow.inherit = False
        tf = s.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, big, 19, col, bold=True, first=True, space=1, align=PP_ALIGN.CENTER)
        para(tf, small, 11, MUTE, space=0, align=PP_ALIGN.CENTER)


def table(slide, rows, x, y, w, col_w, size=11.5, row_h=Inches(0.36), hl=None):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Emu(int(row_h) * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = row_h
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            head_ = i == 0
            c.fill.fore_color.rgb = (INK if head_ else
                                     RGBColor(0xE6, 0xF0, 0xFF) if hl == i else
                                     (PANEL if i % 2 else WHITE))
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            r = p.add_run()
            r.text = val
            r.font.size = Pt(size)
            r.font.name = "Calibri"
            r.font.bold = head_ or j == 0 or hl == i
            r.font.color.rgb = WHITE if head_ else (ACC if (hl == i and j == 0) else INK)
    return t


def set_text(shape, idx, text):
    """Replace paragraph idx's text, keeping its first run's formatting."""
    p = shape.text_frame.paragraphs[idx]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def find(slide, startswith):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(startswith):
            return sh
    raise KeyError(startswith)


def move(slide, new_index):
    lst = prs.slides._sldIdLst
    el = [e for e in lst if prs.slides.get(int(e.get("id"))) is slide] if False else None
    ids = list(lst)
    for e in ids:
        if e.rId == [r for r in ids if prs.part.related_part(r.rId) is slide.part][0].rId:
            lst.remove(e)
            lst.insert(new_index, e)
            return


slides = list(prs.slides)

# ------------------------------------------------------------ stale text fixes
s1 = slides[0]
set_text(find(s1, "Branch 3D_railhead_upgrade"), 0,
         "Branch 3D_railhead_upgrade  ·  status as of 2026-09-24")
sh = find(s1, "A metasurface")
set_text(sh, 1, "Physics ported from the experimentally-validated Face3D codebase, re-derived "
                "at 60 GHz, illuminated by a model of a real 60 GHz horn (RFspin H-A75-W20), "
                "and verified by eleven gates on an RTX 5090 plus a new horn gate (V9).")
s5 = slides[4]
set_text(find(s5, "The RS-I kernel"), 0,
         "The RS-I kernel and the physical-optics currents are unchanged from Face3D. The horn "
         "is now a real 60 GHz part with the textbook aperture model — see the horn slides.")
for sh in s5.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "cosθ/R²" in r.text:
                    r.text = r.text.replace("cosθ/R²", "cosθ/R")
s10 = slides[9]
set_text(find(s10, "Everything chosen in wavelengths"), 0,
         "Everything chosen in wavelengths scales with λ; everything physical (rail, defects, "
         "jitter) stays fixed — and since 2026-09-24 so does the horn, which is now a real part.")
s17 = slides[16]
set_text(find(s17, "Method of Moments"), 0,
         "Solver: Lumerical FDTD, with rail3D's own horn as an Import source (no TFSF). The "
         "setup and the rung ladder are on the next slide.")

# ------------------------------------------------------------ H1: which horn
h1 = prs.slides.add_slide(BLANK)
head(h1, "horn source", "Modelling a real 60 GHz horn: RFspin H-A75-W20",
     "The previous horn was Face3D's Ka-band horn scaled by 5/8 — not a part you can buy, and "
     "its feed carried four waveguide modes at 60 GHz.")
table(h1, [
    ["part", "band", "feed (a × b, mm)", "feed modes at 60 GHz", "verdict"],
    ["previous model", "—", "5.81 × 3.88", "TE10, TE01, TE11, TE20 all propagate", "overmoded, not a real part"],
    ["H-A60-W20", "40–60 GHz", "WR-19  4.78 × 2.39", "single-mode, but 4% below TE20/TE01 (62.8 GHz)", "band edge"],
    ["H-A75-W20", "50–75 GHz", "WR-15  3.76 × 1.88", "TE10 only: cutoff 39.9 GHz, next 79.7 GHz (1.50× cutoff)", "chosen — mid-band"],
    ["H-A90-W20", "60–90 GHz", "WR-12  3.10 × 1.55", "1.24× cutoff: β/k = 0.59, strongly dispersive", "band edge"],
], Inches(0.6), Inches(1.95), Inches(12.13), [1.75, 1.2, 1.9, 4.9, 2.38], size=12, hl=3)
picture(h1, ASSETS / "horn_drawing.png", y=Inches(3.85), h=Inches(2.95))
note(h1, "Inner A, B, L are fitted: RFspin publish the outer shell (23.8 × 17.8 mm, 3D model) and "
         "19–21 dBi over 50–75 GHz; 0.5 mm walls give 19.07 / 20.09 / 21.01 dBi. Confirm with "
         "RFspin's drawing or calipers — wall 0.3–1.0 mm and L 26–31 mm move the rail "
         "illumination < 0.3%.", size=11.5)

# ------------------------------------------------------------ H2: formulas
h2 = prs.slides.add_slide(BLANK)
head(h2, "horn model", "How the horn's amplitude and phase are calculated")
picture(h2, ASSETS / "horn_equations.png", y=Inches(1.5), h=Inches(5.3), x=Inches(0.45),
        w=Inches(7.6))
t = tb(h2, Inches(8.3), Inches(1.6), Inches(4.45), Inches(5.2))
para(t, "Both: an analytic field, then a grid of point sources", 16.5, INK, bold=True,
     first=True, space=8)
para(t, "The aperture distribution is the textbook result for a pyramidal horn (18.38): the "
        "feed's TE10 cosine, stretched to the mouth, times the quadratic phase lag from the "
        "flare's longer path to the edges.", 13, INK, space=8)
para(t, "It is sampled at 20 × 20 cell midpoints, and each sample radiates as a Huygens source "
        "through the same exact RS-I kernel the rail facets use.", 13, INK, space=12)
para(t, "Two fixes vs. Face3D", 16.5, INK, bold=True, space=6)
para(t, "•  Flare phase uses free-space k, not the feed's β_g — with WR-15 at 60 GHz β_g = "
        "0.747k, a 25% phase error.", 13, INK, space=4)
para(t, "•  Midpoint sampling (ΔS = AB/N²); Face3D's edge-inclusive grid over-weighted "
        "amplitude by 2.2%.", 13, INK, space=12)
para(t, "Limit (Nikolova p.15): the aperture method ignores reflections inside the horn and "
        "diffraction at its edges — which is what FDTD rung −1 measures.", 12.5, MUTE,
     italic=True, space=0)
note(h2, "Source: N. K. Nikolova, Lecture 18 — Rectangular Horn Antennas (McMaster Univ.), eqs. "
         "18.4, 18.21, 18.38, 18.39, 18.42–18.44; also Balanis, Antenna Theory, Ch. 13.",
     size=11.5)

# ------------------------------------------------------------ H3: aperture + beam
h3 = prs.slides.add_slide(BLANK)
head(h3, "horn model", "The modelled H-A75-W20: aperture field and beam",
     "Computed by the code the simulator runs (field3d.aperture_distribution), not drawn by hand.")
picture(h3, ASSETS / "horn_aperture.png", y=Inches(1.9), h=Inches(3.75))
stat_row(h3, [("19.1 / 20.1 / 21.0 dBi", "gain at 50 / 60 / 75 GHz — spec 19–21", GOOD),
              ("15.8° / 17.2°", "half-power beamwidth, E / H plane", ACC),
              ("0.53", "aperture efficiency (textbook optimum 0.51)", INK),
              ("TE10 only", "single-mode WR-15 feed at 60 GHz", GOOD),
              ("0.44", "rail distance / far-field 2D²/λ (was 0.73)", INK)],
         y=Inches(5.75), h=Inches(0.95))
note(h3, "V9 gate: aperture-integral directivity (18.21) = closed form (18.39) = 20.09 dBi; gain "
         "inside the part's spec across its band; 20×20 grid converged against 80×80 (corr "
         "0.999998); feed single-mode.", size=11.5, y=Inches(6.88))

# ------------------------------------------------------------ H4: rail
h4 = prs.slides.add_slide(BLANK)
head(h4, "horn model", "What the new horn does on the rail",
     "|E| on a flat plane at crown height, 140 mm from the horn at 55° — the rail sits in the "
     "horn's radiating near field.")
picture(h4, ASSETS / "horn_rail.png", y=Inches(1.95), h=Inches(3.9))
stat_row(h4, [("0.977", "rail illumination, old vs new horn (complex corr)", ACC),
              ("42 → 36 mm", "−3 dB footprint along the rail", INK),
              ("0.28 → 0.32", "|E| at the rail ends, of peak", INK),
              ("≥ 0.997", "across the dimension uncertainty", GOOD)],
         y=Inches(5.9), h=Inches(0.9))
note(h4, "The horn is a provenance key: old datasets are refused by design, and the prelim set "
         "regenerates into a new root.", size=11.5, y=Inches(6.88))

# ------------------------------------------------------------ F1: FDTD setup
f1 = prs.slides.add_slide(BLANK)
head(f1, "full-wave cross-check", "FDTD setup: Lumerical, with the horn as a custom source",
     "No TFSF: rail3D's horn field is written on a plane as an Import source (E and H), so the "
     "rail is lit by the same horn the simulator models.")
table(f1, [
    ["rung", "what runs in Lumerical", "scored by", "pass"],
    ["−1", "horn alone: WR-15 + flare (PEC STL), Mode source TE10 (neff 0.747)",
     "fdtd_agreement.py --horn", "aperture corr ≥ 0.95, D ±0.5 dB, HPBW ±1.5°"],
    ["0", "Import source in an empty box",
     "fdtd_agreement.py --injection", "injected field corr ≥ 0.98"],
    ["1–3", "flat plate → intact rail → cracked rail",
     "fdtd_agreement.py (z = 30 & MS plane)", "per-pixel at z = 30; corr + barcode at MS"],
], Inches(0.6), Inches(1.95), Inches(7.4), [0.55, 2.95, 2.05, 1.85], size=11, row_h=Inches(0.62))
t = tb(f1, Inches(0.6), Inches(4.6), Inches(7.4), Inches(2.2))
para(t, "Four things to get right", 15.5, INK, bold=True, first=True, space=5)
for txt in ["Lumerical is exp(−iωt), like rail3D: exports compare as-is, no conjugation.",
            "The GPU solver has no TFSF; Import sources on GPU need 2025 R1.1 or later.",
            "STL imports as µm unless the length unit is set to mm first.",
            "Material: “PEC (Perfect Electrical Conductor)”. Horn plane z = 15 mm, injecting −z; "
            "the z = 30 monitor then sees only the scattered field."]:
    para(t, "•  " + txt, 12.5, INK, space=3)
mock = ROOT / "data/figures/lumerical_target_setup.png"
if mock.exists():
    picture(f1, mock, y=Inches(1.95), h=Inches(4.7), x=Inches(8.2), w=Inches(4.55))
    t = tb(f1, Inches(8.2), Inches(6.55), Inches(4.55), Inches(0.3))
    para(t, "target layout (mockup, lumerical_mockup.py)", 10.5, MUTE, italic=True, first=True,
         space=0, align=PP_ALIGN.CENTER)
note(f1, "Files: horn_fdtd_case.py (rung −1 STL + plan, λ/20 ≈ 22.5 M cells), horn_source.py "
         "(Import source; --aperture-from uses the FDTD horn instead), compare_wavefronts.py "
         "--export-case (rail). Runbook: LUMERICAL.md.", size=11.5)

# ------------------------------------------------------------ R1: prelim result
r1 = prs.slides.add_slide(BLANK)
head(r1, "first results", "Prelim training run (2026-09-11): pipeline clean, optimiser not",
     "8400 samples in 0.57 h on the 5090, all gates green, 130 → 8 detectors. Run with the "
     "PREVIOUS horn — to be repeated with the H-A75-W20.")
stat_row(r1, [("0.880", "val AUC, trained metasurface", BAD),
              ("0.976", "val AUC, no metasurface", GOOD),
              ("0.000", "SLM at zero phase vs baseline (exact)", INK),
              ("s0, not depth", "crack detection 0.03 → 0.82 across the head", ACC)],
         y=Inches(1.95), h=Inches(1.0))
picture(r1, ASSETS / "prelim_performance.png", y=Inches(3.1), h=Inches(3.55))
note(r1, "The baseline lies inside the metasurface's hypothesis space, so this is an "
         "optimisation failure, not a finding about metasurfaces. Suspects: random-diffuser "
         "initialisation and the capture term; ablate_surface.py tests both.", size=11.5)

# ------------------------------------------------------------ order + save
# horn slides after "the solver" (slide 5), FDTD setup after slide 17's full-wave
# slide, results before the status slide
order = [h1, h2, h3, h4]
for i, s in enumerate(order):
    move(s, 5 + i)
move(f1, 17 + len(order))
move(r1, 18 + len(order))
prs.save(DECK)
print(f"wrote {DECK} ({len(prs.slides._sldIdLst)} slides); backup {BACKUP.name}")
